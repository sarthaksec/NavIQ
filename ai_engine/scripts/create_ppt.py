from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

def create_presentation():
    # Load the template
    template_path = r'd:\programs\sih_isro\SIH2026-IDEA-Presentation-Format.pptx'
    prs = Presentation(template_path)
    
    # Define our content
    ps_id = "SIH1423"
    ps_title = "GNSS-Denied Navigation using Smartphone IMU"
    theme = "Space Technology (ISRO)"
    team_id = "TEAM-ASTRO-99"
    team_name = "AstroNav"

    # Slide 1: Cover
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if hasattr(shape, 'text'):
            if "Problem Statement ID" in shape.text:
                shape.text = f"Problem Statement ID: {ps_id}"
            elif "Problem Statement Title" in shape.text:
                shape.text = f"Problem Statement Title: {ps_title}"
            elif "Theme" in shape.text:
                shape.text = f"Theme: {theme}"
            elif "PS Category" in shape.text:
                shape.text = "PS Category: Software"
            elif "Team ID" in shape.text:
                shape.text = f"Team ID: {team_id}"
            elif "Team Name" in shape.text:
                shape.text = f"Team Name: {team_name}"
    
    # Slide 2: Idea Title
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if hasattr(shape, 'text'):
            if "IDEA TITLE" in shape.text:
                shape.text = "IDEA TITLE: AI-Powered Inertial Navigation Engine"
            elif "Proposed Solution" in shape.text:
                text_frame = shape.text_frame
                text_frame.clear()
                
                p = text_frame.paragraphs[0]
                p.text = "Proposed Solution"
                p.font.bold = True
                p.font.size = Pt(24)
                
                p = text_frame.add_paragraph()
                p.text = " A software engine utilizing Smartphone IMU sensors to maintain lane-level accuracy when GPS is lost."
                p.level = 1
                
                p = text_frame.add_paragraph()
                p.text = " Replaces traditional Dead Reckoning with a Deep Learning velocity estimator."
                p.level = 1
                
                p = text_frame.add_paragraph()
                p.text = "Innovation & Uniqueness"
                p.font.bold = True
                p.font.size = Pt(20)
                
                p = text_frame.add_paragraph()
                p.text = " 1D-CNN Predictor: Filters out vibration to predict precise velocity."
                p.level = 1
                
                p = text_frame.add_paragraph()
                p.text = " HMM Map-Matching: Snaps drifting coordinates to logical road segments."
                p.level = 1

    # Slide 3: Technical Approach
    slide3 = prs.slides[2]
    for shape in slide3.shapes:
        if hasattr(shape, 'text'):
            if "Technologies to be used" in shape.text:
                text_frame = shape.text_frame
                text_frame.clear()
                p = text_frame.paragraphs[0]
                p.text = "Technologies Used: Python, PyTorch, React Three Fiber, ONNX"
                
                p = text_frame.add_paragraph()
                p.text = "Methodology:"
                p.font.bold = True
                
                p = text_frame.add_paragraph()
                p.text = "1. Sensor Ingestion: Read IMU data in real-time."
                p.level = 1
                p = text_frame.add_paragraph()
                p.text = "2. AI Prediction: 1D-CNN outputs continuous velocity vectors."
                p.level = 1
                p = text_frame.add_paragraph()
                p.text = "3. Sensor Fusion: UKF calculates spatial position."
                p.level = 1
                p = text_frame.add_paragraph()
                p.text = "4. Map-Matching: HMM snaps vehicle to the road."
                p.level = 1
                
    # Insert Methodology Image to Slide 3
    img1_path = r'C:\Users\sarth\.gemini\antigravity-ide\brain\764d65f3-fa9a-4279-a465-a77cf4a0457a\methodology_diagram_1788519112908.jpg'
    if os.path.exists(img1_path):
        slide3.shapes.add_picture(img1_path, Inches(5.5), Inches(1.5), width=Inches(4))

    # Slide 4: Feasibility
    slide4 = prs.slides[3]
    for shape in slide4.shapes:
        if hasattr(shape, 'text'):
            if "Analysis of the feasibility" in shape.text:
                text_frame = shape.text_frame
                text_frame.clear()
                
                p = text_frame.paragraphs[0]
                p.text = "Feasibility: High. Uses widely available smartphone sensors and lightweight models that run on edge devices via ONNX."
                
                p = text_frame.add_paragraph()
                p.text = "Challenges:"
                p.font.bold = True
                p = text_frame.add_paragraph()
                p.text = " Unpredictable smartphone orientation."
                p.level = 1
                p = text_frame.add_paragraph()
                p.text = " High battery consumption from continuous inference."
                p.level = 1
                
                p = text_frame.add_paragraph()
                p.text = "Strategies for Overcoming:"
                p.font.bold = True
                p = text_frame.add_paragraph()
                p.text = " Implementing PCA-based gravity alignment for orientation normalization."
                p.level = 1
                p = text_frame.add_paragraph()
                p.text = " Quantizing the PyTorch model to INT8 for efficiency."
                p.level = 1

    # Slide 5: Impact
    slide5 = prs.slides[4]
    for shape in slide5.shapes:
        if hasattr(shape, 'text'):
            if "Potential impact" in shape.text:
                text_frame = shape.text_frame
                text_frame.clear()
                
                p = text_frame.paragraphs[0]
                p.text = "Impact on Target Audience"
                p.font.bold = True
                p = text_frame.add_paragraph()
                p.text = " Drivers in complex infrastructure (tunnels, multi-level highways) will no longer miss exits due to GPS loss."
                p.level = 1
                
                p = text_frame.add_paragraph()
                p.text = "Key Benefits"
                p.font.bold = True
                p = text_frame.add_paragraph()
                p.text = " Economic: Reduces logistical delays for fleet tracking in urban canyons."
                p.level = 1
                p = text_frame.add_paragraph()
                p.text = " Safety: Ensures continuous emergency vehicle tracking."
                p.level = 1
                
    # Insert Impact Image to Slide 5
    img2_path = r'C:\Users\sarth\.gemini\antigravity-ide\brain\764d65f3-fa9a-4279-a465-a77cf4a0457a\impact_diagram_1788519153331.jpg'
    if os.path.exists(img2_path):
        slide5.shapes.add_picture(img2_path, Inches(5.5), Inches(1.5), width=Inches(4))

    # Slide 6: Research
    slide6 = prs.slides[5]
    for shape in slide6.shapes:
        if hasattr(shape, 'text'):
            if "Details / Links" in shape.text:
                text_frame = shape.text_frame
                text_frame.clear()
                p = text_frame.paragraphs[0]
                p.text = "Dataset: Utilized the official IO-VNBD Benchmark Dataset for diverse driver behavior profiling."
                p = text_frame.add_paragraph()
                p.text = "Algorithms Reference:"
                p = text_frame.add_paragraph()
                p.text = " Viterbi Algorithm for Hidden Markov Model (HMM) Map-Matching."
                p.level = 1
                p = text_frame.add_paragraph()
                p.text = " Unscented Kalman Filter mathematics for Non-Linear Sensor Fusion."
                p.level = 1

    # Save
    out_path = r'd:\programs\sih_isro\SIH2026-IDEA-Presentation-Final.pptx'
    prs.save(out_path)
    print(f"Presentation saved to {out_path}")

if __name__ == "__main__":
    create_presentation()
